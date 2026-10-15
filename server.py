import os
import base64
import logging
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from mcp.server.mcpserver import MCPServer


# --------------------------------------------------
# CONFIGURATION
# --------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent

TOKEN_FILE = Path(
    os.environ.get("GMAIL_TOKEN_PATH", str(BASE_DIR / "token.json"))
)
CREDENTIALS_FILE = Path(
    os.environ.get("GMAIL_CREDENTIALS_PATH", str(BASE_DIR / "credentials.json"))
)
WEB_AUTH_MODE = bool(os.environ.get("GMAIL_TOKEN_PATH"))

GENERATED_DIR = BASE_DIR / "generated_files"

SCOPES = [
    "https://www.googleapis.com/auth/gmail.readonly"
]

GENERATED_DIR.mkdir(parents=True, exist_ok=True)


# MCP uses stdout for communication. Logs must go to stderr.
logging.basicConfig(
    level=logging.INFO,
    stream=sys.stderr,
    format="%(asctime)s - %(levelname)s - %(message)s",
)

server = MCPServer("Gmail Agent Server")


# --------------------------------------------------
# GOOGLE OAUTH AUTHENTICATION
# --------------------------------------------------


def get_gmail_credentials():
    """Load and refresh the correct user's Gmail credentials."""

    creds = None

    if TOKEN_FILE.is_file():
        try:
            creds = Credentials.from_authorized_user_file(
                str(TOKEN_FILE),
                SCOPES,
            )
        except (ValueError, OSError) as exc:
            logging.warning("Could not load Gmail token: %s", exc)

            if WEB_AUTH_MODE:
                raise RuntimeError(
                    "The signed-in user's Gmail token is invalid. "
                    "Sign out and authorize Gmail again."
                ) from exc

    if creds and creds.expired and creds.refresh_token:
        try:
            logging.info("Refreshing Gmail token...")
            creds.refresh(Request())

            TOKEN_FILE.write_text(
                creds.to_json(),
                encoding="utf-8",
            )
        except Exception as exc:
            logging.exception("Gmail token refresh failed.")

            if WEB_AUTH_MODE:
                raise RuntimeError(
                    "Could not refresh this user's Gmail token. "
                    "Sign in and authorize Gmail again."
                ) from exc

            creds = None

    if creds and creds.valid:
        return creds

    # A web request must never open the desktop login flow or silently
    # fall back to a different user's token.
    if WEB_AUTH_MODE:
        raise RuntimeError(
            "No valid Gmail token exists for this signed-in user. "
            "Please sign in with Google again."
        )

    # Preserve the original desktop OAuth flow when running server.py
    # directly for standalone development.
    if not CREDENTIALS_FILE.is_file():
        raise FileNotFoundError(
            f"OAuth credentials file not found: {CREDENTIALS_FILE}"
        )

    logging.info("Starting standalone Google account authorization...")

    flow = InstalledAppFlow.from_client_secrets_file(
        str(CREDENTIALS_FILE),
        SCOPES,
    )

    creds = flow.run_local_server(
        port=0,
        prompt="select_account",
        authorization_prompt_message="",
        success_message=(
            "Gmail connected successfully. "
            "You may close this browser tab."
        ),
    )

    TOKEN_FILE.write_text(creds.to_json(), encoding="utf-8")
    return creds


def get_gmail_service():
    """Create an authenticated Gmail API service."""

    creds = get_gmail_credentials()

    return build(
        "gmail",
        "v1",
        credentials=creds,
        cache_discovery=False,
    )


# --------------------------------------------------
# EMAIL HELPERS
# --------------------------------------------------

def decode_email_header(value):
    """Decode an email header."""

    if not value:
        return ""

    from email.header import decode_header, make_header

    try:
        return str(make_header(decode_header(value)))
    except (ValueError, TypeError):
        return value


def decode_message_part(data):
    """Decode URL-safe Base64 email content."""

    if not data:
        return ""

    try:
        data += "=" * (-len(data) % 4)
        decoded = base64.urlsafe_b64decode(data)
        return decoded.decode("utf-8", errors="replace")
    except Exception as exc:
        logging.warning("Could not decode message content: %s", exc)
        return ""


def extract_body(payload):
    """Extract plain text from a Gmail message, with HTML fallback."""

    mime_type = payload.get("mimeType", "")
    body_data = payload.get("body", {}).get("data")

    if mime_type == "text/plain" and body_data:
        return decode_message_part(body_data)

    parts = payload.get("parts", [])

    # Prefer plain-text content.
    for part in parts:
        if part.get("mimeType") == "text/plain":
            data = part.get("body", {}).get("data")
            if data:
                return decode_message_part(data)

    # Handle nested multipart sections.
    for part in parts:
        result = extract_body(part)
        if result:
            return result

    # Use HTML content if plain text is unavailable.
    if mime_type == "text/html" and body_data:
        return decode_message_part(body_data)

    return ""


def get_header_value(headers, name):
    """Get a header value without case sensitivity."""

    for header in headers or []:
        if header.get("name", "").lower() == name.lower():
            return decode_email_header(header.get("value", ""))

    return ""


def format_email(message, include_body=True, body_limit=3000):
    """Format a Gmail message for the AI assistant."""

    payload = message.get("payload", {})
    headers = payload.get("headers", [])

    sender = get_header_value(headers, "From")
    subject = get_header_value(headers, "Subject")
    date = get_header_value(headers, "Date")

    result = [
        f"Message ID: {message.get('id', '')}",
        f"From: {sender or 'Unknown sender'}",
        f"Subject: {subject or '(No subject)'}",
        f"Date: {date or 'Unknown date'}",
    ]

    if include_body:
        body = extract_body(payload)
        if body_limit:
            body = body[:body_limit]

        result.append(f"Body:\n{body or '(No readable text body)'}")

    return "\n".join(result)


def fetch_messages(query, max_results=10):
    """Fetch Gmail messages matching a Gmail search query."""

    service = get_gmail_service()

    response = (
        service.users()
        .messages()
        .list(
            userId="me",
            q=query,
            maxResults=max(1, min(int(max_results), 100)),
        )
        .execute()
    )

    results = []

    for item in response.get("messages", []):
        message = (
            service.users()
            .messages()
            .get(
                userId="me",
                id=item["id"],
                format="full",
            )
            .execute()
        )
        results.append(message)

    return results


# --------------------------------------------------
# TOOL 1: TODAY'S EMAILS
# --------------------------------------------------

@server.tool()
def get_todays_emails():
    """Get today's emails with sender, subject, date, and body."""

    today = datetime.now().strftime("%Y/%m/%d")
    messages = fetch_messages(f"after:{today}", max_results=20)

    if not messages:
        return "No emails found for today."

    return "\n\n" + ("\n\n" + "-" * 50 + "\n\n").join(
        format_email(message, body_limit=3000)
        for message in messages
    )


# --------------------------------------------------
# TOOL 2: UNREAD EMAILS
# --------------------------------------------------

@server.tool()
def get_unread_emails():
    """Get unread emails with sender, subject, date, and body."""

    messages = fetch_messages("is:unread", max_results=10)

    if not messages:
        return "There are no unread emails."

    return "\n\n" + ("\n\n" + "-" * 50 + "\n\n").join(
        format_email(message, body_limit=1500)
        for message in messages
    )


# --------------------------------------------------
# TOOL 3: UNREAD EMAIL COUNT
# --------------------------------------------------

@server.tool()
def get_unread_email_count():
    """Get the estimated number of unread Gmail messages."""

    service = get_gmail_service()

    response = (
        service.users()
        .messages()
        .list(
            userId="me",
            q="is:unread",
            maxResults=1,
        )
        .execute()
    )

    count = response.get("resultSizeEstimate", 0)

    return f"You have approximately {count} unread email(s)."


# --------------------------------------------------
# TOOL 4: SEARCH EMAILS
# --------------------------------------------------

@server.tool()
def search_emails(query: str, max_results: int = 10):
    """
    Search Gmail.

    Examples: interview, subject:job, from:person@example.com,
    newer_than:7d, is:unread.
    """

    if not query or not query.strip():
        return "Please provide a Gmail search query."

    messages = fetch_messages(
        query.strip(),
        max_results=max_results,
    )

    if not messages:
        return f"No emails found for Gmail query: '{query}'."

    return "\n\n" + ("\n\n" + "-" * 50 + "\n\n").join(
        format_email(message, body_limit=2500)
        for message in messages
    )


# --------------------------------------------------
# TOOL 5: GET A SPECIFIC EMAIL
# --------------------------------------------------

@server.tool()
def get_email(message_id: str):
    """Retrieve one email using its Gmail message ID."""

    if not message_id or not message_id.strip():
        return "Please provide a Gmail message ID."

    service = get_gmail_service()

    try:
        message = (
            service.users()
            .messages()
            .get(
                userId="me",
                id=message_id.strip(),
                format="full",
            )
            .execute()
        )

        return format_email(message, body_limit=10000)

    except Exception:
        logging.exception("Could not retrieve Gmail message.")
        return (
            f"Could not retrieve email with ID '{message_id}'. "
            "Check the ID and Gmail authorization."
        )


# --------------------------------------------------
# TOOL 6: CREATE A WORD DOCUMENT
# --------------------------------------------------

@server.tool()
def create_word_file(title: str, content: str):
    """Create a Word document and return its saved path."""

    if not content or not content.strip():
        return "Cannot create a Word document with empty content."

    title = (title or "Generated Document").strip()

    safe_title = "".join(
        ch if ch.isalnum() or ch in (" ", "-", "_") else "_"
        for ch in title
    ).strip()

    safe_title = safe_title[:100] or "Generated Document"

    file_path = GENERATED_DIR / f"{safe_title}.docx"
    counter = 1

    while file_path.exists():
        file_path = GENERATED_DIR / f"{safe_title}_{counter}.docx"
        counter += 1

    document = Document()
    document.add_heading(title, level=1)

    for line in content.splitlines():
        document.add_paragraph(line)

    document.save(str(file_path))

    logging.info("Created Word document: %s", file_path)

    return f"Word document created successfully: {file_path}"


# --------------------------------------------------
# START MCP SERVER
# --------------------------------------------------

if __name__ == "__main__":
    logging.info("Starting Gmail MCP server...")
    server.run()
